import openpyxl
import urllib.request
import json
import re
import time

SERVICE_KEY = 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImxleXdmaXlmbHVpemdraGp4dHhsIiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc4NjkyNzMwNiwiZXhwIjoyMTAyNTAzMzA2fQ.fjUlTXqpI8XIQ3mhb_FkxiUTR093zM3ESV32J0O6z78'
SUPABASE_URL = 'https://leywfiyfluizgkhjxtxl.supabase.co'

def norm_rut(r):
    if not r: return ''
    return re.sub(r'[^0-9kK]', '', str(r)).upper()

def clean_val(v):
    if v is None: return ''
    s = str(v).strip()
    return s

wb = openpyxl.load_workbook('public/Carga datos flota Dasai.xlsx')
ws = wb.active

headers = [clean_val(c.value) for c in ws[1]]
print('Excel headers:', headers)

req = urllib.request.Request(f'{SUPABASE_URL}/rest/v1/tecnicos?select=*', headers={
    'apikey': SERVICE_KEY,
    'Authorization': f'Bearer {SERVICE_KEY}'
})
with urllib.request.urlopen(req) as resp:
    db_tecs = json.loads(resp.read().decode('utf-8'))

db_by_rut = {}
for d in db_tecs:
    dt = d.get('data') or {}
    r = norm_rut(dt.get('rut'))
    if r:
        db_by_rut[r] = d

print(f'Found {len(db_tecs)} existing drivers in DB.')

updated_count = 0
inserted_count = 0

rows = list(ws.iter_rows(min_row=2, values_only=True))
for idx, row in enumerate(rows):
    if not any(row): continue
    row_dict = {headers[i]: clean_val(row[i]) for i in range(min(len(headers), len(row)))}
    
    rut_raw = row_dict.get('RUT', '')
    rut_clean = norm_rut(rut_raw)
    name = row_dict.get('CONDUCTOR', '')
    ppu = row_dict.get('PPU', '').upper()
    phone = row_dict.get('CELULAR', '')
    phone2 = row_dict.get('WHATSAPP', '')
    email = row_dict.get('MAIL CONDUCTOR', '')
    gps_cccs = row_dict.get('GPS CCCS', '')
    beetrack = row_dict.get('BEETRACK', '')
    induccion = row_dict.get('INDUCCION', '')
    carpeta = row_dict.get('CARPETA', '')
    contrato = row_dict.get('CONTRATO', '')
    anexo = row_dict.get('ANEXO', '')
    
    dueno = ''
    for k in row_dict:
        if 'DUE' in k.upper():
            dueno = row_dict[k]
            break
            
    tipo_vehic = row_dict.get('TIPO VEHIC', '')
    facturacion = row_dict.get('FACTURACION', '')
    nombre_empresa = row_dict.get('NOMBRE EMPRESA', '')
    rut_empresa = row_dict.get('RUT EMPRESA', '')
    banco = row_dict.get('BANCO', '')
    tipo_cuenta = row_dict.get('TIPO CUENTA', '')
    numero_cuenta = row_dict.get('NUMERO CUENTA', '')
    comuna = row_dict.get('COMUNA', '')
    
    direccion = ''
    for k in row_dict:
        if 'DIREC' in k.upper():
            direccion = row_dict[k]
            break
            
    licencia = row_dict.get('LICENCIA', '')
    gps = row_dict.get('GPS', '')
    seguro = row_dict.get('SEGURO', '')

    existing = db_by_rut.get(rut_clean)
    if existing:
        existing_data = existing.get('data') or {}
        merged_data = {
            **existing_data,
            'name': name or existing_data.get('name', ''),
            'rut': rut_raw or existing_data.get('rut', ''),
            'patente': ppu or existing_data.get('patente', ''),
            'phone': phone or existing_data.get('phone', ''),
            'phone2': phone2 or existing_data.get('phone2', ''),
            'email': email or existing_data.get('email', ''),
            'gpsCccs': gps_cccs,
            'beetrack': beetrack,
            'induccion': induccion,
            'carpeta': carpeta,
            'contrato': contrato,
            'anexo': anexo,
            'duenoFurgon': dueno,
            'tipoVehiculo': tipo_vehic,
            'facturacion': facturacion,
            'nombreEmpresa': nombre_empresa,
            'rutEmpresa': rut_empresa,
            'banco': banco,
            'tipoCuenta': tipo_cuenta,
            'numeroCuenta': numero_cuenta,
            'comuna': comuna or existing_data.get('comuna', ''),
            'direccion': direccion or existing_data.get('direccion', ''),
            'licencia': licencia,
            'gps': gps,
            'seguro': seguro,
            'tipoServicio': existing_data.get('tipoServicio') or 'Paquetería',
        }
        ex_id = existing['id']
        update_url = f'{SUPABASE_URL}/rest/v1/tecnicos?id=eq.{ex_id}'
        req_up = urllib.request.Request(update_url, data=json.dumps({'data': merged_data}).encode('utf-8'), headers={
            'apikey': SERVICE_KEY,
            'Authorization': f'Bearer {SERVICE_KEY}',
            'Content-Type': 'application/json',
            'Prefer': 'return=representation'
        }, method='PATCH')
        with urllib.request.urlopen(req_up) as r_up:
            pass
        updated_count += 1
        print(f'Updated [{ex_id}] {name}')
    else:
        new_id = f'tech-{int(time.time() * 1000) + idx}'
        new_data = {
            'id': new_id,
            'name': name,
            'rut': rut_raw,
            'patente': ppu,
            'phone': phone,
            'phone2': phone2,
            'email': email,
            'gpsCccs': gps_cccs,
            'beetrack': beetrack,
            'induccion': induccion,
            'carpeta': carpeta,
            'contrato': contrato,
            'anexo': anexo,
            'duenoFurgon': dueno,
            'tipoVehiculo': tipo_vehic,
            'facturacion': facturacion,
            'nombreEmpresa': nombre_empresa,
            'rutEmpresa': rut_empresa,
            'banco': banco,
            'tipoCuenta': tipo_cuenta,
            'numeroCuenta': numero_cuenta,
            'comuna': comuna,
            'direccion': direccion,
            'licencia': licencia,
            'gps': gps,
            'seguro': seguro,
            'status': 'disponible',
            'tipoServicio': 'Paquetería',
            'completedOrders': 0,
            'avgTime': 0,
            'productivity': 0,
            'documentos': {},
            'autoDocumentos': {}
        }
        insert_url = f'{SUPABASE_URL}/rest/v1/tecnicos'
        req_in = urllib.request.Request(insert_url, data=json.dumps({'id': new_id, 'data': new_data, 'tech_number': len(db_tecs) + inserted_count + 1}).encode('utf-8'), headers={
            'apikey': SERVICE_KEY,
            'Authorization': f'Bearer {SERVICE_KEY}',
            'Content-Type': 'application/json',
            'Prefer': 'return=representation'
        }, method='POST')
        with urllib.request.urlopen(req_in) as r_in:
            pass
        inserted_count += 1
        print(f'Inserted [{new_id}] {name}')

print(f'\n--- Summary ---')
print(f'Total in Excel: {len(rows)}')
print(f'Updated: {updated_count}')
print(f'Inserted: {inserted_count}')
